"""素材导入通道服务：前端统一入口，把用户上传的结构化素材做成待采纳草案。

信任模型与经营基线草案一致（draft→gate→adopt）：
  1. 上传：原件算 sha256 存 imports/drafts/<id>/raw/，只读不可变；建 ImportJob 行(uploaded)。
  2. 分析：md 先经 vfab_markdown_to_tables.py 扁平化成 CSV；模型在【已声明本体类清单】约束下
     为每个文件建议 target_class / entity_keys / time_field，产出 manifest 草案(analyzed)。
  3. 校验：把草案 manifest 落到 staging 目录，跑 imported_ingest.py --check 真门禁(validated/invalid)。
  4. 采纳（仅人工）：合入线上 sources/internal/imported/manifest.json + 拷原件 → 跑线上门禁 →
     失败回滚、同名冲突拒绝(adopted)。restricted 素材禁止采纳。

绝不自动写入推理层/current.ttl。模型只产出草案建议，一切过引擎脚本真门禁，人点采纳才落线上。
"""
from __future__ import annotations

import hashlib
import json
import re
import shutil
import uuid
from datetime import datetime, timezone
from pathlib import Path

from sqlalchemy.orm import Session

from ..config import settings
from ..models import ImportJob, User
from .llm import ExternalServiceError, llm_service, user_api_key, user_llm_endpoint
from .orchestrator import _json_object
from .semi_kb import SemiKbError, semi_kb

JOB_ID_RE = re.compile(r"^[0-9]+-[0-9a-f]{12}$")
_DRAFTS_ROOT = "imports/drafts"
_LIVE_SOURCE = "sources/internal/imported"

# 结构化直接入库的格式；md 需先扁平化成 csv，txt 只暂存不自动入库。
STRUCTURED_FORMATS = {"xlsx", "xls", "csv", "tsv", "json"}
TABLEIFY_FORMATS = {"md", "markdown"}
STAGE_ONLY_FORMATS = {"txt"}
ALLOWED_UPLOAD_FORMATS = STRUCTURED_FORMATS | TABLEIFY_FORMATS | STAGE_ONLY_FORMATS
ALLOWED_CLASSIFICATIONS = {"internal", "internal_confidential", "restricted"}
MAX_UPLOAD_BYTES = 25 * 1024 * 1024  # 单文件 25MB 上限，防止上传超大文件占满磁盘


def _job_dir(job_id: str) -> Path:
    if not JOB_ID_RE.match(job_id or ""):
        raise SemiKbError(f"非法导入任务 ID：{job_id}")
    return settings.engine_root / "imports" / "drafts" / job_id


def job_dir_path(job_id: str) -> Path:
    job_dir = _job_dir(job_id)
    if not job_dir.is_dir():
        raise SemiKbError(f"导入任务目录不存在：{job_id}")
    return job_dir


def _ext_format(filename: str) -> str:
    ext = Path(filename or "").suffix.lower().lstrip(".")
    return "markdown" if ext == "md" else ext


def _safe_stem(value: str) -> str:
    return re.sub(r"[^a-zA-Z0-9._-]+", "-", value).strip("-.")[:80] or "file"


async def create_import_job(
    db: Session, user: User, files: list[tuple[str, bytes]], classification: str
) -> dict:
    """存原件（sha256 锁定）+ 建 ImportJob 行。files 为 [(filename, content_bytes)]。"""
    if classification not in ALLOWED_CLASSIFICATIONS:
        raise ValueError(f"非法 classification：{classification}")
    if not files:
        raise ValueError("未提供任何文件")

    job_id = f"{user.id}-{uuid.uuid4().hex[:12]}"
    job_dir = _job_dir(job_id)
    raw_dir = job_dir / "raw"
    raw_dir.mkdir(parents=True, exist_ok=True)

    stored: list[dict] = []
    seen_names: set[str] = set()
    for filename, content in files:
        fmt = _ext_format(filename)
        if fmt not in ALLOWED_UPLOAD_FORMATS:
            shutil.rmtree(job_dir, ignore_errors=True)
            raise ValueError(f"不支持的格式：{filename}（仅 xlsx/xls/csv/tsv/json/md/txt）")
        if len(content) > MAX_UPLOAD_BYTES:
            shutil.rmtree(job_dir, ignore_errors=True)
            raise ValueError(f"文件过大（>25MB）：{filename}")
        stem = _safe_stem(Path(filename).stem)
        base = f"{stem}.{'md' if fmt == 'markdown' else fmt}"
        name = base
        suffix = 1
        while name in seen_names:  # 同名去重，避免覆盖
            name = f"{stem}-{suffix}.{'md' if fmt == 'markdown' else fmt}"
            suffix += 1
        seen_names.add(name)
        (raw_dir / name).write_bytes(content)
        stored.append({
            "original_name": filename, "stored_name": name, "format": fmt,
            "sha256": hashlib.sha256(content).hexdigest(), "size_bytes": len(content),
        })

    meta = {
        "job_id": job_id, "user_id": user.id, "classification": classification,
        "uploaded": stored, "created_at": datetime.now(timezone.utc).isoformat(),
    }
    (job_dir / "meta.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")

    row = ImportJob(
        id=job_id, user_id=user.id, status="uploaded", classification=classification,
        summary=f"上传 {len(stored)} 个文件", manifest_json=json.dumps({"files": []}, ensure_ascii=False),
    )
    db.add(row)
    db.commit()
    return job_summary(row, uploaded=stored)


def job_summary(row: ImportJob, uploaded: list[dict] | None = None) -> dict:
    manifest = json.loads(row.manifest_json or "{}")
    payload = {
        "job_id": row.id, "status": row.status, "channel": row.channel,
        "classification": row.classification, "summary": row.summary,
        "llm_model_id": row.llm_model_id, "manifest": manifest,
        "validation": json.loads(row.validation_json or "{}"),
        "adopted_paths": json.loads(row.adopted_paths_json or "{}"),
        "created_at": row.created_at, "adopted_at": row.adopted_at, "rejected_at": row.rejected_at,
    }
    if uploaded is not None:
        payload["uploaded"] = uploaded
    return payload


def _read_headers(path: Path, fmt: str) -> list[str]:
    """读结构化文件首行表头（后端侧，供模型建议映射；与引擎门禁的 file_headers 同义）。"""
    import csv as _csv

    if fmt in {"csv", "tsv"}:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return [h.strip() for h in next(_csv.reader(handle, delimiter="\t" if fmt == "tsv" else ","), []) if h.strip()]
    if fmt == "json":
        data = json.loads(path.read_text(encoding="utf-8"))
        row = data[0] if isinstance(data, list) and data else data
        return [str(k) for k in row] if isinstance(row, dict) else []
    if fmt in {"xlsx", "xls"}:
        try:
            from openpyxl import load_workbook
        except ModuleNotFoundError:
            return []
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        first = next(sheet.iter_rows(max_row=1, values_only=True), ())
        workbook.close()
        return [str(cell).strip() for cell in first if cell is not None and str(cell).strip()]
    return []


async def _tableify_markdown(md_path: Path, out_dir: Path) -> list[Path]:
    """把 md 经引擎 vfab_markdown_to_tables.py 扁平化成 CSV（不 in-process import 引擎脚本）。"""
    out_dir.mkdir(parents=True, exist_ok=True)
    result = await semi_kb.command(
        "vfab_markdown_to_tables.py", str(md_path), "--out-dir", str(out_dir), timeout=120
    )
    if result["exit_code"]:
        raise SemiKbError(f"markdown 扁平化失败：{result['output'][-500:]}")
    return sorted(out_dir.glob("*.csv"))


def _declared_classes(limit: int = 800) -> list[dict]:
    """已声明本体类清单（{iri,label}），用于约束模型建议的 target_class。"""
    context = semi_kb.ontology_context(limit=limit)
    return [{"iri": t["iri"], "label": t.get("label") or t["iri"]}
            for t in context.get("terms") or [] if t.get("kind") == "class"]


def _build_mapping_prompt(files: list[dict], classes: list[dict]) -> tuple[str, str]:
    class_lines = "\n".join(f"  - {c['iri']}（{c['label']}）" for c in classes[:400])
    system = (
        "你是半导体知识库的素材导入映射助手。给定若干结构化数据文件的表头，为每个文件建议它最贴切的"
        "本体目标类(target_class)、实体主键(entity_keys)与时间字段(time_field)。只输出一个 JSON 对象，"
        "不要任何解释或 Markdown 代码块。JSON 结构：\n"
        '{"summary": "一句话概括这批素材", "files": [{"stored_name": "原样回填", '
        '"target_class": "必须从下方已声明类清单里选一个 IRI", "entity_keys": ["从该文件表头里选的主键列名"], '
        '"time_field": "该文件表头里的时间列名或 null"}]}\n\n'
        "硬规则：\n"
        "1. target_class 必须严格等于下方清单中的某个 IRI，禁止杜撰或改写；拿不准就选语义最接近的。\n"
        "2. entity_keys 每一项都必须是该文件 headers 里真实存在的列名（区分大小写）。\n"
        "3. time_field 必须是该文件 headers 里的列名或 null，不要臆造。\n"
        "4. 每个输入文件都要在 files 里给出一条，stored_name 原样回填。\n\n"
        "已声明本体类清单（只能从中选 target_class）：\n" + class_lines
    )
    user = json.dumps(
        {"files": [{"stored_name": f["stored_name"], "headers": f["headers"], "format": f["format"]} for f in files]},
        ensure_ascii=False,
    )
    return system, user


async def analyze_import_job(db: Session, user: User, job_id: str, model_id: str) -> dict:
    """扁平化 md → 读表头 → 模型在已声明类约束下建议映射 → 写 staging manifest 草案。"""
    row = db.get(ImportJob, job_id)
    if not row or row.user_id != user.id:
        raise SemiKbError(f"导入任务不存在：{job_id}")
    api_key = user_api_key(db, user.id)
    if not api_key:
        raise ExternalServiceError("未配置模型 API Key，无法分析导入素材")

    job_dir = job_dir_path(job_id)
    meta = json.loads((job_dir / "meta.json").read_text(encoding="utf-8"))
    staging_data = job_dir / "staging" / "raw"
    if staging_data.exists():
        shutil.rmtree(staging_data, ignore_errors=True)
    staging_data.mkdir(parents=True, exist_ok=True)

    prepared: list[dict] = []  # 待入库的结构化文件（结构化原件 + md 派生 csv）
    staged_only: list[dict] = []  # txt 等仅暂存、不自动入库
    for entry in meta["uploaded"]:
        fmt = entry["format"]
        src = job_dir / "raw" / entry["stored_name"]
        if fmt in STAGE_ONLY_FORMATS:
            staged_only.append({"stored_name": entry["stored_name"], "format": fmt,
                                "note": "散文/纯文本仅暂存，需人工转结构化后再导入"})
            continue
        if fmt in TABLEIFY_FORMATS:
            csvs = await _tableify_markdown(src, staging_data / Path(entry["stored_name"]).stem)
            for csv_path in csvs:
                dest = staging_data / f"{Path(entry['stored_name']).stem}__{csv_path.name}"
                if csv_path.resolve() != dest.resolve():
                    shutil.move(str(csv_path), dest)
                prepared.append({"stored_name": dest.name, "format": "csv",
                                 "headers": _read_headers(dest, "csv"), "derived_from": entry["stored_name"]})
        else:
            dest = staging_data / entry["stored_name"]
            shutil.copy2(src, dest)
            prepared.append({"stored_name": entry["stored_name"], "format": fmt,
                             "headers": _read_headers(dest, fmt), "derived_from": None})

    return await _suggest_and_persist(db, row, model_id, api_key, prepared, staged_only)


async def _suggest_and_persist(
    db: Session, row: ImportJob, model_id: str, api_key: str,
    prepared: list[dict], staged_only: list[dict],
) -> dict:
    classes = _declared_classes()
    valid_iris = {c["iri"] for c in classes}
    suggestions: dict[str, dict] = {}
    summary = row.summary
    if prepared:
        system, user_prompt = _build_mapping_prompt(prepared, classes)
        raw = await llm_service.complete(
            api_key, model_id, system, user_prompt, temperature=0.1,
            endpoint=user_llm_endpoint(db, row.user_id),
        )
        parsed = _json_object(raw)
        summary = str(parsed.get("summary") or summary)[:1000]
        for item in parsed.get("files") or []:
            if isinstance(item, dict) and item.get("stored_name"):
                suggestions[str(item["stored_name"])] = item

    files: list[dict] = []
    for prep in prepared:
        headers = prep["headers"]
        hint = suggestions.get(prep["stored_name"], {})
        target = str(hint.get("target_class") or "")
        keys = [k for k in (hint.get("entity_keys") or []) if k in headers]
        time_field = hint.get("time_field") if hint.get("time_field") in headers else None
        files.append({
            "stored_name": prep["stored_name"], "format": prep["format"], "headers": headers,
            "derived_from": prep.get("derived_from"),
            "target_class": target if target in valid_iris else "",  # 未命中已声明类留空待人工选
            "entity_keys": keys, "time_field": time_field,
            "target_class_valid": target in valid_iris,
        })

    manifest = {"files": files, "staged_only": staged_only, "declared_classes": classes}
    row.manifest_json = json.dumps(manifest, ensure_ascii=False)
    row.llm_model_id = model_id
    row.summary = summary
    row.status = "analyzed"
    row.validation_json = "{}"
    db.commit()
    return job_summary(row)


def update_mapping(db: Session, user: User, job_id: str, files: list[dict]) -> dict:
    """人工修订每文件的 target_class/entity_keys/time_field。只改映射，不动原件。"""
    row = db.get(ImportJob, job_id)
    if not row or row.user_id != user.id:
        raise SemiKbError(f"导入任务不存在：{job_id}")
    manifest = json.loads(row.manifest_json or "{}")
    by_name = {f["stored_name"]: f for f in manifest.get("files") or []}
    valid_iris = {c["iri"] for c in manifest.get("declared_classes") or []}
    for edit in files:
        name = str(edit.get("stored_name") or "")
        target = by_name.get(name)
        if not target:
            continue
        headers = target.get("headers") or []
        if "target_class" in edit:
            tc = str(edit["target_class"] or "")
            target["target_class"] = tc
            target["target_class_valid"] = tc in valid_iris
        if "entity_keys" in edit:
            target["entity_keys"] = [k for k in (edit.get("entity_keys") or []) if k in headers]
        if "time_field" in edit:
            tf = edit.get("time_field")
            target["time_field"] = tf if tf in headers else None
    row.manifest_json = json.dumps(manifest, ensure_ascii=False)
    row.status = "analyzed"  # 改过映射需重新校验
    row.validation_json = "{}"
    db.commit()
    return job_summary(row)


def _writable_files(manifest: dict) -> list[dict]:
    """可入库文件：已选中已声明的 target_class 的结构化文件。"""
    return [f for f in manifest.get("files") or [] if f.get("target_class") and f.get("target_class_valid")]


def _build_staging_manifest(row: ImportJob, manifest: dict) -> dict:
    return {
        "source_id": "internal.imported",
        "delivered_at": datetime.now(timezone.utc).isoformat(),
        "owner": f"user:{row.user_id}",
        "classification": row.classification,
        "files": [{"path": f"raw/{f['stored_name']}", "format": f["format"],
                   "sha256": hashlib.sha256((_job_dir(row.id) / "staging" / "raw" / f["stored_name"]).read_bytes()).hexdigest(),
                   "target_class": f["target_class"], "entity_keys": f.get("entity_keys") or [],
                   "time_field": f.get("time_field")}
                  for f in _writable_files(manifest)],
    }


async def validate_import_job(db: Session, user: User, job_id: str) -> dict:
    """把草案 manifest 落到 staging 目录，跑 imported_ingest.py --check 真门禁（不碰线上）。"""
    row = db.get(ImportJob, job_id)
    if not row or row.user_id != user.id:
        raise SemiKbError(f"导入任务不存在：{job_id}")
    if row.classification == "restricted":
        raise SemiKbError("restricted 素材禁止入库校验；只可暂存与记录元数据")
    manifest = json.loads(row.manifest_json or "{}")
    if not _writable_files(manifest):
        raise SemiKbError("没有可入库的文件：请为至少一个文件选择已声明的目标类")

    job_dir = job_dir_path(job_id)
    staging = job_dir / "staging"
    (staging / "manifest.json").write_text(
        json.dumps(_build_staging_manifest(row, manifest), ensure_ascii=False, indent=2), encoding="utf-8"
    )
    rel = staging.relative_to(settings.engine_root).as_posix()
    result = await semi_kb.command("imported_ingest.py", "--source-dir", rel, "--check", timeout=300)
    payload = semi_kb._json_from_output(result["output"])
    passed = payload.get("status") == "available"
    validation = {"passed": passed, "status": payload.get("status"),
                  "error": payload.get("error"), "datasets": payload.get("datasets") or [],
                  "duration_seconds": result["duration_seconds"]}
    row.validation_json = json.dumps(validation, ensure_ascii=False)
    row.status = "validated" if passed else "invalid"
    db.commit()
    return job_summary(row)


async def adopt_import_job(db: Session, user: User, job_id: str) -> dict:
    """人点『采纳』：合入线上 sources/internal/imported/ → 跑线上门禁 → 失败回滚、同名冲突拒绝。"""
    row = db.get(ImportJob, job_id)
    if not row or row.user_id != user.id:
        raise SemiKbError(f"导入任务不存在：{job_id}")
    if row.status == "adopted":
        return job_summary(row)
    if row.classification == "restricted":
        raise SemiKbError("restricted 素材禁止采纳入库")
    validation = json.loads(row.validation_json or "{}")
    if not validation.get("passed"):
        raise SemiKbError("草案未通过引擎校验，无法采纳；请先校验通过")

    manifest = json.loads(row.manifest_json or "{}")
    job_dir = job_dir_path(job_id)
    staging_raw = job_dir / "staging" / "raw"
    async with semi_kb._candidate_lock:  # 与经营基线晋升同一把锁，串行写线上源目录
        live = settings.engine_root / "sources" / "internal" / "imported"
        live_raw = live / "raw"
        live_raw.mkdir(parents=True, exist_ok=True)
        live_manifest_path = live / "manifest.json"
        manifest_backup = live_manifest_path.read_bytes() if live_manifest_path.is_file() else None
        existing = json.loads(manifest_backup.decode("utf-8")) if manifest_backup else \
            {"source_id": "internal.imported", "owner": "material-import", "classification": "internal_confidential", "files": []}
        existing_paths = {f.get("path") for f in existing.get("files") or []}

        new_entries: list[dict] = []
        written: list[Path] = []
        try:
            for f in _writable_files(manifest):
                src = staging_raw / f["stored_name"]
                dest_name = f"{row.id}__{f['stored_name']}"
                rel_path = f"raw/{dest_name}"
                if rel_path in existing_paths or (live_raw / dest_name).exists():
                    raise SemiKbError(f"线上已存在同名素材文件：{rel_path}")
                shutil.copy2(src, live_raw / dest_name)
                written.append(live_raw / dest_name)
                new_entries.append({"path": rel_path, "format": f["format"],
                                    "sha256": hashlib.sha256(src.read_bytes()).hexdigest(),
                                    "target_class": f["target_class"], "entity_keys": f.get("entity_keys") or [],
                                    "time_field": f.get("time_field")})
            merged = {**existing, "source_id": "internal.imported",
                      "delivered_at": datetime.now(timezone.utc).isoformat(),
                      "classification": existing.get("classification") or "internal_confidential",
                      "files": (existing.get("files") or []) + new_entries}
            live_manifest_path.write_text(json.dumps(merged, ensure_ascii=False, indent=2), encoding="utf-8")
            result = await semi_kb.command("imported_ingest.py", "--source-dir", "sources/internal/imported", timeout=600)
            if result["exit_code"]:
                raise SemiKbError("采纳后线上门禁失败：" + result["output"][-2000:])
        except BaseException:
            for path in written:
                path.unlink(missing_ok=True)
            if manifest_backup is not None:
                live_manifest_path.write_bytes(manifest_backup)  # 复原原 manifest
            else:
                live_manifest_path.unlink(missing_ok=True)  # 本次首建则删除
            raise
        semi_kb.invalidate_cache()

    row.status = "adopted"
    row.adopted_at = datetime.now(timezone.utc)
    row.adopted_paths_json = json.dumps({"files": [e["path"] for e in new_entries],
                                         "gate_output": result["output"][-1000:]}, ensure_ascii=False)
    db.commit()
    discard_job_files(job_id)
    return job_summary(row)


def discard_job_files(job_id: str) -> None:
    """删除导入任务 staging 目录（幂等）。原件与 staging 都在其中。"""
    job_dir = _job_dir(job_id)
    if job_dir.is_dir():
        shutil.rmtree(job_dir, ignore_errors=True)
